import os
import requests
from bs4 import BeautifulSoup
import pandas as pd
from openpyxl import load_workbook
from io import BytesIO
from urllib.parse import urljoin
import re

def setup_directories(base_dir):
    """Set up base directory."""
    os.makedirs(base_dir, exist_ok=True)

def fetch_html_content(url):
    """Fetch page content with error handling."""
    response = requests.get(url)
    response.raise_for_status()
    return response.content

def extract_excel_links(soup):
    """Find all Excel file links on the page."""
    return [a['href'] for a in soup.find_all("a", href=lambda href: href and href.endswith(".xlsx"))]

def download_excel_file(session, file_url):
    """Download Excel file using a requests session."""
    response = session.get(file_url, stream=True)
    response.raise_for_status()
    return BytesIO(response.content)

def save_csv_from_sheet(data, folder_path, month, year, sheet_name):
    """Process and save each Excel sheet as CSV."""
    try:
        columns = next(data)
        df = pd.DataFrame(data, columns=columns)
        clean_sheet_name = re.sub(r'\W+', '_', sheet_name).strip('_')
        csv_path = os.path.join(folder_path, f"{month}-{year}-{clean_sheet_name}.csv")
        df.to_csv(csv_path, index=False)
        print(f"  Saved: {csv_path}")
    except StopIteration:
        print(f"  Sheet '{sheet_name}' is empty. Skipping.")

def fetch_eurocom_data():
    base_dir = "Data/INFORM Suite/INFORM Severity/EuroCom_INFORM_Severity_Data"
    setup_directories(base_dir)

    base_url = "https://drmkc.jrc.ec.europa.eu"
    page_url = "https://drmkc.jrc.ec.europa.eu/inform-index/INFORM-Severity/Results-and-data"
    
    # Fetch and parse HTML content
    html_content = fetch_html_content(page_url)
    soup = BeautifulSoup(html_content, "html.parser")
    
    # Extract Excel file links
    excel_links = extract_excel_links(soup)
    if not excel_links:
        print("No Excel links found.")
        return

    print(f"Found {len(excel_links)} Excel file to download.")
    
    # Session for optimized network requests
    with requests.Session() as session:
        for idx, relative_url in enumerate(excel_links, start=1):
            file_url = urljoin(base_url, relative_url)
            file_name = os.path.basename(relative_url)
            
            # Extract month and year from filename
            match = re.search(r"_(\w+)_(\d{4})", file_name)
            if match:
                month, year = match.groups()
                folder_path = os.path.join(base_dir, f"{year} {month}")
                os.makedirs(folder_path, exist_ok=True)

                try:
                    print(f"[{idx}/{len(excel_links)}] Downloading {file_name}...")
                    excel_data = download_excel_file(session, file_url)
                    
                    # Process each sheet
                    workbook = load_workbook(excel_data, data_only=True)
                    for sheet_name in workbook.sheetnames:
                        sheet = workbook[sheet_name]
                        data = sheet.values
                        save_csv_from_sheet(data, folder_path, month, year, sheet_name)
                
                except requests.HTTPError as e:
                    print(f"Failed to download {file_name}: {e}")
                except Exception as e:
                    print(f"Error processing {file_name}: {e}")

    print("All files downloaded.")


